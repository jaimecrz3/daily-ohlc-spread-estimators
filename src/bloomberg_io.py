"""Read-only Bloomberg ingestion, preserving cached errors and physical rows.

Only values-only sheets supply numerical/reference values. Formula workbooks
are inspected solely for exact field order, security references and overrides.
No recalculation, workbook saves, imputation or duplicate resolution.
"""
from pathlib import Path
from datetime import datetime, date
import hashlib
import re

import numpy as np
import openpyxl
import pandas as pd

DAILY_FIELDS = ['PX_OPEN', 'PX_HIGH', 'PX_LOW', 'PX_LAST', 'PX_VOLUME', 'PX_BID', 'PX_ASK']
REF_FIELDS = ['NAME', 'ID_BB_GLOBAL', 'ID_ISIN', 'ID_CUSIP', 'TICKER_AND_EXCH_CODE', 'CRNCY']
SIDE_FIELDS = ['SECURITY_TYP', 'MARKET_SECTOR_DES', 'EQY_PRIM_EXCH_SHRT']
OHLC = DAILY_FIELDS[:4]
KEYS = ['asset', 'date']


def sha256(path):
    digest = hashlib.sha256()
    with Path(path).open('rb') as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b''):
            digest.update(block)
    return digest.hexdigest()


def verify_hashes(raw, manifest_path=None):
    """Verify safe relative paths, sizes and SHA before any workbook reads."""
    raw = Path(raw).resolve()
    path = Path(manifest_path) if manifest_path else raw / 'sha256_manifest.csv'
    manifest = pd.read_csv(path, dtype={'file_name': str, 'sha256': str})
    if list(manifest.columns) != ['file_name', 'sha256', 'size_bytes'] or manifest.file_name.duplicated().any():
        raise ValueError('Invalid or duplicate hash manifest records')
    rows = []
    for entry in manifest.itertuples(index=False):
        target = (raw / entry.file_name).resolve()
        if Path(entry.file_name).is_absolute() or not target.is_relative_to(raw) or target == path.resolve():
            raise ValueError(f'Unsafe manifest path: {entry.file_name}')
        exists = target.is_file()
        actual = sha256(target) if exists else None
        size = target.stat().st_size if exists else None
        rows.append(dict(file_name=entry.file_name, expected_sha256=entry.sha256,
            actual_sha256=actual, expected_size_bytes=entry.size_bytes,
            actual_size_bytes=size, hash_matches=actual == entry.sha256,
            size_matches=size == entry.size_bytes, exists=exists))
    audit = pd.DataFrame(rows)
    return audit, manifest


def require_verified(audit):
    if audit.empty or not (audit.hash_matches & audit.size_matches & audit.exists).all():
        raise ValueError('Raw acquisition hash verification failed; stop before ingestion')


def formula_text(value):
    return getattr(value, 'text', value)


def inspect_daily_lineage(path, selected, convention):
    """Verify actual formula anchors, never use formula caches for prices."""
    workbook = openpyxl.load_workbook(path, data_only=False, read_only=True)
    rows = []
    try:
        for security in selected.itertuples():
            sheet = workbook[f'{security.asset} US EQUITY']
            expected = f'{security.asset} US Equity'
            if sheet['A1'].value != expected or sheet['B1'].value != 20150101 or sheet['C1'].value != 20251231:
                raise ValueError(f'Unexpected request security/window: {sheet.title}')
            formula = formula_text(sheet['A4'].value)
            if not isinstance(formula, str) or 'BDH(' not in formula:
                raise ValueError('Expected saved BDH formula anchor at A4')
            quoted = re.findall(r'"([^"]*)"', formula)
            if not quoted or quoted[0].split(',') != DAILY_FIELDS:
                raise ValueError('BDH field ordering differs from expected column mapping')
            overrides = {part.split('=', 1)[0]: part.split('=', 1)[1] for part in quoted[1:] if '=' in part}
            expected_flags = {'CapChg': 'T', 'CshAdjNormal': 'T'} if convention == 'adjusted' else {'CapChg': 'F', 'CshAdjNormal': 'F', 'CshAdjAbnormal': 'F'}
            if any(overrides.get(k) != v for k, v in expected_flags.items()):
                raise ValueError(f'Unexpected adjustment overrides: {formula}')
            rows.append(dict(asset=security.asset, convention=convention,
                security=expected, sheet=sheet.title, formula_cell='A4', formula=formula,
                requested_fields=','.join(DAILY_FIELDS), overrides=';'.join(f'{k}={v}' for k,v in sorted(overrides.items())),
                special_cash_dividend_override=overrides.get('CshAdjAbnormal', 'not explicitly specified'),
                source_file=Path(path).name))
    finally:
        workbook.close()
    return pd.DataFrame(rows)


def read_reference(values_path, formula_path, selected):
    values = openpyxl.load_workbook(values_path, data_only=False, read_only=True)
    formulas = openpyxl.load_workbook(formula_path, data_only=False, read_only=True)
    records, comparison = [], []
    try:
        sheet, lineage = values['Reference'], formulas['Reference']
        if [str(sheet.cell(1, c).value).strip() for c in range(2, 8)] != REF_FIELDS:
            raise ValueError('Unexpected reference header mapping')
        if [sheet.cell(6, c).value for c in range(9, 12)] != SIDE_FIELDS:
            raise ValueError('Unexpected side-block layout')
        lookup = selected.set_index('asset')
        observed = []
        for row in range(2, 20):
            security = sheet.cell(row, 1).value
            asset = security.split()[0]
            observed.append(asset)
            if asset not in lookup.index:
                raise ValueError(f'Unexpected reference asset: {asset}')
            record = dict(asset=asset, asset_class=lookup.loc[asset, 'asset_class'],
                liquidity_tier=lookup.loc[asset, 'liquidity_tier'], bloomberg_security=security,
                source_file=Path(values_path).name, worksheet='Reference', main_source_row=row,
                side_source_row=row+5, reference_snapshot_date='2026-10-13')
            for col, field in enumerate(REF_FIELDS, 2):
                formula = formula_text(lineage.cell(row, col).value)
                if not isinstance(formula, str) or f'$A{row},"{field}"' not in formula:
                    raise ValueError('Main reference field/security lineage mismatch')
                value = sheet.cell(row, col).value
                if sheet.cell(row, col).data_type == 'f':
                    raise ValueError('Values-only reference workbook contains formula')
                record[field] = value
            for col, field in enumerate(SIDE_FIELDS, 9):
                formula = formula_text(lineage.cell(row+5, col).value)
                if not isinstance(formula, str) or f'$A{row},"{field}"' not in formula:
                    raise ValueError('Side-block field/security lineage mismatch')
                if sheet.cell(row+5, col).data_type == 'f':
                    raise ValueError('Values-only side block contains formula')
                record[field] = sheet.cell(row+5, col).value
            for field in REF_FIELDS + SIDE_FIELDS:
                record[field+'_error_flag'] = isinstance(record[field], str) and record[field].startswith('#')
                record[field+'_missing_flag'] = record[field] is None
            cls, typ = record['asset_class'], record['SECURITY_TYP']
            record['instrument_type_consistent'] = typ in ({'ETP', 'ETF'} if cls == 'etfs' else {'Common Stock', 'REIT', 'ADR', 'ADS'})
            record['currency_usd_confirmed'] = record['CRNCY'] == 'USD'
            record['market_sector_equity_confirmed'] = record['MARKET_SECTOR_DES'] == 'Equity'
            record['historical_identity_status'] = 'OCX_to_IMDX_stable_ID_unresolved; no manual stitching' if asset == 'IMDX' else 'current_reference_validated; full_historical_identity_not_certified'
            record['validation_status'] = 'current_type_currency_and_sector_consistent' if all(record[k] for k in ['instrument_type_consistent', 'currency_usd_confirmed', 'market_sector_equity_confirmed']) else 'review_required'
            records.append(record)
        if len(set(observed)) != 18 or set(observed) != set(selected.asset):
            raise ValueError('Reference universe differs from frozen 18 assets')
        for row in (2, 3):
            for col, field in enumerate(['bloomberg_security', 'NAME', 'ID_BB_GLOBAL', 'ID_ISIN', 'TICKER_AND_EXCH_CODE'], 9):
                cell = sheet.cell(row, col)
                comparison.append(dict(comparison_security=sheet.cell(row, 9).value,
                    field=field, cached_value=cell.value, error_flag=isinstance(cell.value, str) and cell.value.startswith('#'),
                    worksheet='Reference', cell=f'{openpyxl.utils.get_column_letter(col)}{row}', source_file=Path(values_path).name))
    finally:
        values.close(); formulas.close()
    return pd.DataFrame(records), pd.DataFrame(comparison)


def normalize_daily_rows(rows, asset, cls, convention, source_file, worksheet):
    """Preserve dated/partial records, blank/error tokens and source row IDs."""
    records = []
    for excel_row, cells in rows:
        cells = list(cells) + [None] * max(0, 8-len(cells))
        if all(v is None for v in cells):
            continue  # Empty formatted worksheet padding is not a returned row.
        raw_date = cells[0]
        valid_date = isinstance(raw_date, (datetime, date, pd.Timestamp))
        record = dict(asset=asset, asset_class=cls, convention=convention,
            date=pd.Timestamp(raw_date).normalize() if valid_date else pd.NaT,
            date_raw=str(raw_date) if raw_date is not None else None,
            unassignable_date=not valid_date, date_has_time=bool(valid_date and pd.Timestamp(raw_date) != pd.Timestamp(raw_date).normalize()),
            source='Bloomberg', source_role='daily_source_validation',
            source_file=source_file, worksheet=worksheet, source_row=excel_row)
        for field, value in zip(DAILY_FIELDS, cells[1:8]):
            raw = str(value) if value is not None else None
            is_number = isinstance(value, (int, float, np.number)) and not isinstance(value, bool)
            number = float(value) if is_number else np.nan
            record[field] = number
            record[field+'_raw'] = raw
            record[field+'_blank'] = value is None
            record[field+'_error'] = isinstance(value, str) and value.startswith('#')
            record[field+'_non_numeric'] = value is not None and not is_number and not record[field+'_error']
            record[field+'_non_finite'] = is_number and not np.isfinite(number)
        records.append(record)
    return annotate_daily(pd.DataFrame(records))


def annotate_daily(data):
    data = data.copy()
    numeric = data[DAILY_FIELDS]
    finite = pd.DataFrame(np.isfinite(numeric), index=data.index, columns=DAILY_FIELDS)
    data['complete_ohlc'] = finite[OHLC].all(axis=1)
    data['complete_ohlcv'] = finite[OHLC+['PX_VOLUME']].all(axis=1)
    data['duplicate_date'] = data.date.notna() & data.duplicated(['asset', 'date'], keep=False)
    data['date_ordering_problem'] = False
    for _, group in data.groupby('asset', sort=False):
        dated = group.loc[group.date.notna()]
        data.loc[dated.index, 'date_ordering_problem'] = dated.date.diff().lt(pd.Timedelta(0))
    data['invalid_ohlc'] = ((data.PX_HIGH < data.PX_LOW)
        | (data.PX_OPEN > data.PX_HIGH) | (data.PX_OPEN < data.PX_LOW)
        | (data.PX_LAST > data.PX_HIGH) | (data.PX_LAST < data.PX_LOW)
        | data[OHLC].le(0).any(axis=1))
    data['nonpositive_price'] = data[[f for f in DAILY_FIELDS if f != 'PX_VOLUME']].le(0).any(axis=1)
    data['negative_volume'] = data.PX_VOLUME.lt(0)
    data['ask_less_than_bid'] = data.PX_ASK.lt(data.PX_BID)
    data['partial_ohlc_row'] = ~data.complete_ohlc
    data['daily_input_valid'] = data.complete_ohlcv & ~data.invalid_ohlc & ~data.negative_volume & ~data.duplicate_date & ~data.unassignable_date
    return data


def read_daily_values(path, selected, convention):
    workbook = openpyxl.load_workbook(path, data_only=False, read_only=True)
    pieces, excluded = [], []
    try:
        expected = {f'{row.asset} US EQUITY': row for row in selected.itertuples()}
        for name, security in expected.items():
            sheet = workbook[name]
            if sheet['A1'].value != f'{security.asset} US Equity':
                raise ValueError('Values sheet security does not match selected asset')
            rows = []
            for number, row in enumerate(sheet.iter_rows(min_row=4, max_col=8), 4):
                if any(c.data_type == 'f' for c in row):
                    raise ValueError('Values-only historical sheet contains formula')
                rows.append((number, [c.value for c in row]))
            pieces.append(normalize_daily_rows(rows, security.asset, security.asset_class,
                convention, Path(path).name, name))
        for name in workbook.sheetnames:
            if name not in expected:
                for row in workbook[name]:
                    for cell in row:
                        if cell.value is not None:
                            excluded.append(dict(convention=convention, source_file=Path(path).name,
                                worksheet=name, cell=cell.coordinate, cached_value=str(cell.value),
                                role='excluded_diagnostic_sheet_not_final_universe'))
    finally:
        workbook.close()
    return pd.concat(pieces, ignore_index=True), pd.DataFrame(excluded)


def unique_dated(data):
    dated = data.loc[data.date.notna()].copy()
    if dated.duplicated(KEYS).any():
        raise ValueError('Duplicate asset/date keys: comparison cannot use many-to-many joins')
    return dated


def overlap_join(left, right):
    return unique_dated(left).merge(unique_dated(right), on=KEYS,
        how='inner', suffixes=('_left', '_right'), validate='one_to_one', sort=True)


def summarize_daily(data):
    rows = []
    for asset, group in data.groupby('asset', sort=False):
        record = dict(asset=asset, convention=group.convention.iloc[0],
            first_returned_date=group.date.min(), last_returned_date=group.date.max(),
            total_returned_rows=len(group), dated_rows=int(group.date.notna().sum()),
            complete_ohlc_days=int(group.complete_ohlc.sum()), complete_ohlcv_days=int(group.complete_ohlcv.sum()),
            duplicate_date_rows=int(group.duplicate_date.sum()),
            duplicate_dates=int(group.loc[group.duplicate_date, 'date'].nunique()),
            invalid_ohlc_rows=int(group.invalid_ohlc.sum()),
            ask_less_than_bid_rows=int(group.ask_less_than_bid.sum()),
            nonpositive_price_rows=int(group.nonpositive_price.sum()),
            negative_volume_rows=int(group.negative_volume.sum()),
            date_ordering_problems=int(group.date_ordering_problem.sum()),
            unassignable_date_rows=int(group.unassignable_date.sum()),
            date_has_time_rows=int(group.date_has_time.sum()))
        for field in DAILY_FIELDS:
            record[field+'_missing_count'] = int((~np.isfinite(group[field])).sum())
            for kind in ['blank', 'error', 'non_numeric', 'non_finite']:
                record[field+'_'+kind+'_count'] = int(group[field+'_'+kind].sum())
        rows.append(record)
    return pd.DataFrame(rows)

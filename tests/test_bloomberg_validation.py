from pathlib import Path
import sys
import tempfile
import unittest

import numpy as np
import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from bloomberg_io import (sha256, verify_hashes, require_verified, normalize_daily_rows,
    read_daily_values, read_reference, inspect_daily_lineage, overlap_join, summarize_daily, DAILY_FIELDS, REF_FIELDS, SIDE_FIELDS)
from bloomberg_validation import compare_conventions, compare_local


def daily(rows, asset='TEST', convention='unadjusted'):
    return normalize_daily_rows([(i+4, [pd.Timestamp(day)]+list(values)) for i,(day,values) in enumerate(rows)],
        asset, 'stocks', convention, f'{convention}.xlsx', 'TEST US EQUITY')


class BloombergValidationTests(unittest.TestCase):
    def test_nested_manifest_paths_verified_without_location_fallback(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            folder = root / '13-10-2026'
            folder.mkdir()
            path = folder / 'reference.txt'
            path.write_text('immutable acquisition fixture')
            digest = sha256(path)
            pd.DataFrame([dict(file_name='13-10-2026/reference.txt',
                sha256=digest, size_bytes=path.stat().st_size)]).to_csv(
                    root / 'sha256_manifest.csv', index=False)
            audit, manifest = verify_hashes(root)
            require_verified(audit)
            self.assertEqual(manifest.file_name.iloc[0], '13-10-2026/reference.txt')
            self.assertEqual(sha256(path), digest)
            # A stale path must fail even when the basename exists elsewhere.
            manifest['file_name'] = 'reference.txt'
            manifest.to_csv(root / 'sha256_manifest.csv', index=False)
            audit, _ = verify_hashes(root)
            with self.assertRaises(ValueError):
                require_verified(audit)

    def test_reference_side_security_mapping_errors_and_immutable_files(self):
        with tempfile.TemporaryDirectory() as directory:
            folder=Path(directory)
            formula=openpyxl.Workbook(); values=openpyxl.Workbook()
            f,v=formula.active,values.active
            f.title=v.title='Reference'
            selected=pd.DataFrame([dict(asset=f'T{i}', asset_class='stocks', liquidity_tier='LOW') for i in range(18)])
            for col,field in enumerate(REF_FIELDS,2):
                v.cell(1,col,field)
            for col,field in enumerate(SIDE_FIELDS,9):
                v.cell(6,col,field)
            for row,asset in enumerate(selected.asset,2):
                v.cell(row,1,asset+' US Equity')
                for col,field in enumerate(REF_FIELDS,2):
                    v.cell(row,col,'USD' if field=='CRNCY' else field+' fixture')
                    f.cell(row,col,'=BDP($A'+str(row)+',"'+field+'")')
                for col,(field,value) in enumerate(zip(SIDE_FIELDS,['Common Stock','Equity','UW']),9):
                    v.cell(row+5,col,value)
                    f.cell(row+5,col,'=BDP($A'+str(row)+',"'+field+'")')
            v['I2']='OCX US Equity'; v['K2']='#N/A Field Not Applicable'; v['L2']='#N/A N/A'
            fp,vp=folder/'formula.xlsx',folder/'values.xlsx'
            formula.save(fp); values.save(vp)
            before=[sha256(fp),sha256(vp)]
            audit,comparison=read_reference(vp,fp,selected)
            self.assertEqual(len(audit),18)
            self.assertTrue(audit.instrument_type_consistent.all())
            self.assertEqual(audit.side_source_row.iloc[-1],24)
            self.assertEqual(comparison.error_flag.sum(),2)
            self.assertEqual(before,[sha256(fp),sha256(vp)])

    def test_hash_success_failure_and_stop(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / 'raw.txt'
            path.write_text('immutable')
            pd.DataFrame([dict(file_name='raw.txt', sha256=sha256(path), size_bytes=path.stat().st_size)]).to_csv(root/'sha256_manifest.csv', index=False)
            audit, _ = verify_hashes(root)
            require_verified(audit)
            path.write_text('tampered fixture')
            audit, _ = verify_hashes(root)
            with self.assertRaises(ValueError):
                require_verified(audit)

    def test_hash_path_escape_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            pd.DataFrame([dict(file_name='../escape', sha256='0'*64, size_bytes=0)]).to_csv(root/'sha256_manifest.csv', index=False)
            with self.assertRaises(ValueError):
                verify_hashes(root)

    def test_missing_error_blank_and_non_numeric_preserved(self):
        result = daily([('2024-01-02', [10, 11, 9, '#N/A N/A', None, 'unknown', 12]),
                        ('2024-01-03', [10, 11, 9, 10, 0, 9, 11])])
        self.assertTrue(np.isnan(result.PX_LAST.iloc[0]))
        self.assertEqual(result.PX_LAST_raw.iloc[0], '#N/A N/A')
        self.assertTrue(result.PX_LAST_error.iloc[0])
        self.assertTrue(result.PX_VOLUME_blank.iloc[0])
        self.assertTrue(result.PX_BID_non_numeric.iloc[0])
        self.assertFalse(result.complete_ohlc.iloc[0])
        self.assertEqual(result.PX_VOLUME.iloc[1], 0)
        self.assertTrue(result.complete_ohlcv.iloc[1])
        self.assertEqual(len(result), 2)

    def test_duplicates_and_order_preserved_and_joins_blocked(self):
        result = daily([('2024-01-03', [10, 11, 9, 10, 1, 9, 11]),
                        ('2024-01-02', [10, 11, 9, 10, 1, 9, 11]),
                        ('2024-01-02', [10, 11, 9, 10, 1, 9, 11])])
        self.assertEqual(result.duplicate_date.sum(), 2)
        self.assertEqual(result.date_ordering_problem.sum(), 1)
        self.assertEqual(len(result), 3)
        with self.assertRaises(ValueError):
            overlap_join(result, result)

    def test_invalid_ohlc_nonpositive_volume_crossed_quote(self):
        result = daily([('2024-01-02', [12, 11, 9, 10, -1, 12, 11]),
                        ('2024-01-03', [10, 11, -1, 10, 1, 0, 11])])
        self.assertTrue(result.invalid_ohlc.all())
        self.assertTrue(result.negative_volume.iloc[0])
        self.assertTrue(result.ask_less_than_bid.iloc[0])
        self.assertTrue(result.nonpositive_price.iloc[1])

    def test_overlap_only_no_calendar_or_imputation(self):
        bloom = daily([('2024-01-02', [10, 11, 9, 10, 2, 9, 11]),
                       ('2024-01-04', [10, 11, 9, None, 2, 9, 11])])
        local = pd.DataFrame(dict(asset=['TEST', 'TEST'], date=pd.to_datetime(['2024-01-04', '2024-01-05']),
            open=[10,10], high=[11,11], low=[9,9], close=[10,10], volume=[2,2], aggregation_valid=[True,True]))
        metrics, coverage, _ = compare_local(bloom, local)
        self.assertEqual(coverage.overlapping_dates.iloc[0], 1)
        self.assertEqual(coverage.bloomberg_only_dates.iloc[0], 1)
        self.assertEqual(coverage.first_rate_only_dates.iloc[0], 1)
        self.assertEqual(metrics.loc[metrics.field.eq('close'), 'valid_numeric_pairs'].iloc[0], 0)
        self.assertTrue(np.isnan(bloom.PX_LAST.iloc[1]))

    def test_adjustment_label_and_volume_differences(self):
        unadj = daily([('2022-07-25', [1,2,.5,1,100,1,2])])
        adj = daily([('2022-07-25', [10,20,5,10,10,10,20])], convention='adjusted')
        summary, _ = compare_conventions(unadj, adj)
        self.assertEqual(summary.ohlc_differing_days.iloc[0], 1)
        self.assertEqual(summary.volume_differing_days.iloc[0], 1)
        self.assertEqual(summary.PX_LAST_max_relative_difference.iloc[0], 9)
        self.assertEqual(adj.convention.iloc[0], 'adjusted')

    def test_workbook_parse_no_raw_modification_padding_excluded(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)/'values.xlsx'
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = 'TEST US EQUITY'
            sheet['A1'] = 'TEST US Equity'
            for col,value in enumerate([pd.Timestamp('2024-01-02').to_pydatetime(),10,11,9,10,1,'#N/A N/A',11],1):
                sheet.cell(4,col,value)
            sheet.cell(20,1).number_format = 'yyyy-mm-dd'
            workbook.create_sheet('OCX_check')['A4'] = '#N/A N/A'
            workbook.save(path)
            before = sha256(path)
            selected = pd.DataFrame([dict(asset='TEST',asset_class='stocks')])
            result, excluded = read_daily_values(path,selected,'unadjusted')
            self.assertEqual(len(result),1)
            self.assertEqual(result.source_row.iloc[0],4)
            self.assertEqual(result.PX_BID_raw.iloc[0],'#N/A N/A')
            self.assertEqual(excluded.worksheet.iloc[0],'OCX_check')
            self.assertEqual(sha256(path),before)

    def test_formula_fields_overrides_verified_without_recalc(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)/'formula.xlsx'
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = 'TEST US EQUITY'
            sheet['A1']='TEST US Equity'; sheet['B1']=20150101; sheet['C1']=20251231
            sheet['A4']='=BDH($A$1,"'+','.join(DAILY_FIELDS)+'",$B$1,$C$1,"CapChg=F","CshAdjNormal=F","CshAdjAbnormal=F")'
            workbook.save(path)
            before=sha256(path)
            selected=pd.DataFrame([dict(asset='TEST')])
            lineage=inspect_daily_lineage(path,selected,'unadjusted')
            self.assertEqual(lineage.special_cash_dividend_override.iloc[0],'F')
            with self.assertRaises(ValueError):
                inspect_daily_lineage(path,selected,'adjusted')
            self.assertEqual(sha256(path),before)

    def test_summary_reports_missing_and_duplicates_separately(self):
        result=daily([('2024-01-02',[10,11,9,10,1,None,'#N/A N/A'])])
        summary=summarize_daily(result).iloc[0]
        self.assertEqual(summary.complete_ohlcv_days,1)
        self.assertEqual(summary.PX_BID_missing_count,1)
        self.assertEqual(summary.PX_ASK_error_count,1)

    def test_zero_reference_relative_difference_is_undefined(self):
        bloom=daily([('2024-01-02',[10,11,9,10,0,9,11])])
        local=pd.DataFrame([dict(asset='TEST',date=pd.Timestamp('2024-01-02'),open=10,high=11,low=9,close=10,volume=0,aggregation_valid=True)])
        metrics,_,_=compare_local(bloom,local)
        volume=metrics.loc[metrics.field.eq('volume')].iloc[0]
        self.assertEqual(volume.valid_numeric_pairs,1)
        self.assertEqual(volume.relative_difference_pairs,0)
        self.assertEqual(volume.exact_match_rate,1)

    def test_deterministic_output_and_no_input_mutation(self):
        bloom=daily([('2024-01-02',[10,11,9,10,2,9,11])])
        original=bloom.copy(deep=True)
        local=pd.DataFrame([dict(asset='TEST',date=pd.Timestamp('2024-01-02'),open=9,high=11,low=9,close=10,volume=1,aggregation_valid=True)])
        first=compare_local(bloom,local)
        second=compare_local(bloom,local)
        for a,b in zip(first,second):
            pd.testing.assert_frame_equal(a,b)
        pd.testing.assert_frame_equal(bloom,original)


if __name__=='__main__':
    unittest.main()
